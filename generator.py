"""
ParseDeck - Presentation Deck Generator
Programmatically transforms extracted academic research paper metadata (JSON / Pydantic)
into a modern, executive-grade PowerPoint presentation using python-pptx.
"""

from __future__ import annotations

import json
import logging
from pathlib import Path
from typing import List, Optional, Union

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from extractor import BenchmarkResult, KeyComponent, PaperAnalysis

logger = logging.getLogger("parsedeck.generator")


# ==============================================================================
# Design System & Palette Constants (16:9 Widescreen)
# ==============================================================================

SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 7.5

FONT_HEADING = "Calibri"
FONT_BODY = "Calibri"

# Colors
COLOR_DARK_BG = RGBColor(15, 23, 42)       # Midnight Navy (#0F172A)
COLOR_DARK_CARD = RGBColor(30, 41, 59)     # Slate Navy (#1E293B)
COLOR_LIGHT_BG = RGBColor(248, 250, 252)   # Off-White (#F8FAFC)
COLOR_LIGHT_CARD = RGBColor(255, 255, 255) # Pure White (#FFFFFF)
COLOR_BORDER = RGBColor(226, 232, 240)     # Subtle Slate Border (#E2E8F0)
COLOR_DARK_BORDER = RGBColor(51, 65, 85)   # Border for Dark Cards (#334155)

# Accent Colors
COLOR_ACCENT_INDIGO = RGBColor(79, 70, 229) # Indigo (#4F46E5)
COLOR_ACCENT_CYAN = RGBColor(14, 165, 233)   # Sky Cyan (#0EA5E9)
COLOR_ACCENT_GREEN = RGBColor(16, 185, 129)  # Emerald (#10B981)
COLOR_ACCENT_AMBER = RGBColor(245, 158, 11)  # Amber (#F59E0B)

# Text Colors
COLOR_TEXT_PRIMARY = RGBColor(15, 23, 42)    # Dark Slate
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)   # Slate Muted
COLOR_TEXT_LIGHT = RGBColor(255, 255, 255)   # White
COLOR_TEXT_LIGHT_MUTED = RGBColor(203, 213, 225) # Light Slate


class PresentationDeckBuilder:
    """Builder class that compiles PaperAnalysis into a multi-slide presentation."""

    def __init__(self, analysis: PaperAnalysis):
        self.analysis = analysis
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_WIDTH_IN)
        self.prs.slide_height = Inches(SLIDE_HEIGHT_IN)
        self.blank_layout = self.prs.slide_layouts[6]  # 100% blank slide layout

    # --------------------------------------------------------------------------
    # Low-level UI Helper Methods
    # --------------------------------------------------------------------------

    def _set_slide_background(self, slide, color: RGBColor):
        """Fills the entire slide with a solid background color."""
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(SLIDE_WIDTH_IN), Inches(SLIDE_HEIGHT_IN)
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = color
        bg_shape.line.fill.background()  # No border

    def _add_header(self, slide, kicker: str, title: str):
        """Adds a standard header banner with kicker and prominent title."""
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Kicker (Category / Section tracker)
        p_kicker = tf.paragraphs[0]
        p_kicker.text = kicker.upper()
        p_kicker.font.name = FONT_HEADING
        p_kicker.font.size = Pt(10)
        p_kicker.font.bold = True
        p_kicker.font.color.rgb = COLOR_ACCENT_INDIGO
        p_kicker.space_after = Pt(2)

        # Slide Main Title
        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.name = FONT_HEADING
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_PRIMARY

    def _add_card(
        self,
        slide,
        left: float,
        top: float,
        width: float,
        height: float,
        fill_color: RGBColor = COLOR_LIGHT_CARD,
        border_color: Optional[RGBColor] = COLOR_BORDER,
    ):
        """Creates a styled card container."""
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
        return card

    # --------------------------------------------------------------------------
    # Slide 1: Title & Executive Summary Slide (Dark Theme)
    # --------------------------------------------------------------------------

    def build_title_slide(self):
        """Builds an executive dark-themed title slide."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_slide_background(slide, COLOR_DARK_BG)

        # Accent Top Badge
        badge = slide.shapes.add_textbox(Inches(0.9), Inches(0.8), Inches(11.5), Inches(0.4))
        tf_b = badge.text_frame
        tf_b.word_wrap = True
        p_b = tf_b.paragraphs[0]
        venue_text = f" // {self.analysis.publication_venue}" if self.analysis.publication_venue else ""
        p_b.text = f"PARSEDECK RESEARCH BRIEFING{venue_text}".upper()
        p_b.font.name = FONT_HEADING
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_ACCENT_CYAN

        # Paper Title
        title_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.3), Inches(11.5), Inches(1.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = self.analysis.title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(28)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_LIGHT

        # Authors
        if self.analysis.authors:
            p_auth = tf_t.add_paragraph()
            authors_str = ", ".join(self.analysis.authors[:6])
            p_auth.text = f"Authors: {authors_str}"
            p_auth.font.name = FONT_BODY
            p_auth.font.size = Pt(13)
            p_auth.font.color.rgb = COLOR_TEXT_LIGHT_MUTED
            p_auth.space_before = Pt(8)

        # Executive Summary Container Card
        self._add_card(
            slide,
            left=0.9,
            top=3.5,
            width=11.5,
            height=3.1,
            fill_color=COLOR_DARK_CARD,
            border_color=COLOR_DARK_BORDER,
        )

        # Left accent stripe on executive summary card
        stripe = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(3.5), Inches(0.12), Inches(3.1)
        )
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = COLOR_ACCENT_INDIGO
        stripe.line.fill.background()

        # Text inside Executive Summary Card
        summary_box = slide.shapes.add_textbox(Inches(1.3), Inches(3.7), Inches(10.8), Inches(2.7))
        tf_s = summary_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0

        p_s_label = tf_s.paragraphs[0]
        p_s_label.text = "EXECUTIVE SUMMARY & CORE THESIS"
        p_s_label.font.name = FONT_HEADING
        p_s_label.font.size = Pt(11)
        p_s_label.font.bold = True
        p_s_label.font.color.rgb = COLOR_ACCENT_CYAN
        p_s_label.space_after = Pt(8)

        p_s_body = tf_s.add_paragraph()
        p_s_body.text = self.analysis.executive_summary
        p_s_body.font.name = FONT_BODY
        p_s_body.font.size = Pt(14)
        p_s_body.font.color.rgb = COLOR_TEXT_LIGHT
        p_s_body.line_spacing = 1.25

    # --------------------------------------------------------------------------
    # Slide 2: Core Problem & Existing Bottlenecks (Light Theme)
    # --------------------------------------------------------------------------

    def build_problem_slide(self):
        """Builds slide detailing the problem statement and existing limitations."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_slide_background(slide, COLOR_LIGHT_BG)
        self._add_header(
            slide,
            kicker="01 // Research Motivation & Challenges",
            title="The Core Problem & Architecture Bottlenecks",
        )

        card_width = 5.65
        card_height = 5.2
        top_pos = 1.7

        # Left Card: The Fundamental Problem
        self._add_card(slide, left=0.8, top=top_pos, width=card_width, height=card_height)
        tb_left = slide.shapes.add_textbox(Inches(1.1), Inches(top_pos + 0.3), Inches(5.05), Inches(4.6))
        tf_l = tb_left.text_frame
        tf_l.word_wrap = True

        p_l_label = tf_l.paragraphs[0]
        p_l_label.text = "FUNDAMENTAL BOTTLENECK"
        p_l_label.font.name = FONT_HEADING
        p_l_label.font.size = Pt(11)
        p_l_label.font.bold = True
        p_l_label.font.color.rgb = COLOR_ACCENT_AMBER
        p_l_label.space_after = Pt(10)

        p_l_body = tf_l.add_paragraph()
        p_l_body.text = self.analysis.core_problem
        p_l_body.font.name = FONT_BODY
        p_l_body.font.size = Pt(13)
        p_l_body.font.color.rgb = COLOR_TEXT_PRIMARY
        p_l_body.line_spacing = 1.25

        # Right Card: Existing System Limitations
        self._add_card(slide, left=6.85, top=top_pos, width=card_width, height=card_height)
        tb_right = slide.shapes.add_textbox(Inches(7.15), Inches(top_pos + 0.3), Inches(5.05), Inches(4.6))
        tf_r = tb_right.text_frame
        tf_r.word_wrap = True

        p_r_label = tf_r.paragraphs[0]
        p_r_label.text = "EXISTING SYSTEM LIMITATIONS"
        p_r_label.font.name = FONT_HEADING
        p_r_label.font.size = Pt(11)
        p_r_label.font.bold = True
        p_r_label.font.color.rgb = COLOR_ACCENT_AMBER
        p_r_label.space_after = Pt(10)

        for limitation in self.analysis.problem_limitations[:4]:
            p_item = tf_r.add_paragraph()
            p_item.text = f"•  {limitation}"
            p_item.font.name = FONT_BODY
            p_item.font.size = Pt(12)
            p_item.font.color.rgb = COLOR_TEXT_PRIMARY
            p_item.space_after = Pt(10)
            p_item.line_spacing = 1.2

    # --------------------------------------------------------------------------
    # Slide 3: Key Technical Innovation (Light Theme)
    # --------------------------------------------------------------------------

    def build_innovation_slide(self):
        """Builds slide showcasing the primary breakthrough and key components."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_slide_background(slide, COLOR_LIGHT_BG)
        self._add_header(
            slide,
            kicker="02 // Architectural Innovation",
            title="Core Technical Breakthrough & Architecture",
        )

        # Top Hero Card: Primary Innovation
        self._add_card(slide, left=0.8, top=1.7, width=11.7, height=1.7, fill_color=COLOR_LIGHT_CARD)
        
        # Left accent stripe on Hero Card
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.7), Inches(0.1), Inches(1.7))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = COLOR_ACCENT_INDIGO
        stripe.line.fill.background()

        tb_hero = slide.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(11.1), Inches(1.4))
        tf_h = tb_hero.text_frame
        tf_h.word_wrap = True

        p_h_tag = tf_h.paragraphs[0]
        p_h_tag.text = "PRIMARY BREAKTHROUGH / PARADIGM SHIFT"
        p_h_tag.font.name = FONT_HEADING
        p_h_tag.font.size = Pt(10)
        p_h_tag.font.bold = True
        p_h_tag.font.color.rgb = COLOR_ACCENT_INDIGO
        p_h_tag.space_after = Pt(4)

        p_h_text = tf_h.add_paragraph()
        p_h_text.text = self.analysis.key_innovation
        p_h_text.font.name = FONT_BODY
        p_h_text.font.size = Pt(13)
        p_h_text.font.color.rgb = COLOR_TEXT_PRIMARY
        p_h_text.line_spacing = 1.2

        # Bottom 3 Component Cards
        components = self.analysis.key_components[:3]
        num_cards = max(len(components), 1)
        total_width = 11.7
        gap = 0.3
        col_width = (total_width - (gap * (num_cards - 1))) / num_cards
        card_top = 3.65
        card_height = 3.25

        for idx, comp in enumerate(components):
            c_left = 0.8 + idx * (col_width + gap)
            self._add_card(slide, left=c_left, top=card_top, width=col_width, height=card_height)

            tb_comp = slide.shapes.add_textbox(
                Inches(c_left + 0.25), Inches(card_top + 0.25), Inches(col_width - 0.5), Inches(card_height - 0.5)
            )
            tf_c = tb_comp.text_frame
            tf_c.word_wrap = True

            p_num = tf_c.paragraphs[0]
            p_num.text = f"COMPONENT 0{idx + 1}"
            p_num.font.name = FONT_HEADING
            p_num.font.size = Pt(9)
            p_num.font.bold = True
            p_num.font.color.rgb = COLOR_ACCENT_CYAN
            p_num.space_after = Pt(4)

            p_name = tf_c.add_paragraph()
            p_name.text = comp.name
            p_name.font.name = FONT_HEADING
            p_name.font.size = Pt(13)
            p_name.font.bold = True
            p_name.font.color.rgb = COLOR_TEXT_PRIMARY
            p_name.space_after = Pt(8)

            p_desc = tf_c.add_paragraph()
            p_desc.text = comp.description
            p_desc.font.name = FONT_BODY
            p_desc.font.size = Pt(11)
            p_desc.font.color.rgb = COLOR_TEXT_MUTED
            p_desc.line_spacing = 1.2

    # --------------------------------------------------------------------------
    # Slide 4: Algorithm Overview & Execution Flow (Light Theme)
    # --------------------------------------------------------------------------

    def build_algorithm_slide(self):
        """Builds slide presenting the algorithm overview and sequential steps."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_slide_background(slide, COLOR_LIGHT_BG)
        self._add_header(
            slide,
            kicker="03 // Methodology & Execution Flow",
            title="Algorithm Overview & Pipeline Mechanics",
        )

        # Overview Banner Card
        self._add_card(slide, left=0.8, top=1.7, width=11.7, height=1.3, fill_color=COLOR_LIGHT_CARD)
        tb_ov = slide.shapes.add_textbox(Inches(1.05), Inches(1.85), Inches(11.2), Inches(1.0))
        tf_o = tb_ov.text_frame
        tf_o.word_wrap = True

        p_o_label = tf_o.paragraphs[0]
        p_o_label.text = "PIPELINE ARCHITECTURE OVERVIEW"
        p_o_label.font.name = FONT_HEADING
        p_o_label.font.size = Pt(10)
        p_o_label.font.bold = True
        p_o_label.font.color.rgb = COLOR_ACCENT_INDIGO
        p_o_label.space_after = Pt(2)

        p_o_body = tf_o.add_paragraph()
        p_o_body.text = self.analysis.algorithm_overview
        p_o_body.font.name = FONT_BODY
        p_o_body.font.size = Pt(12)
        p_o_body.font.color.rgb = COLOR_TEXT_PRIMARY

        # Sequential Step Cards (Horizontal Process Layout)
        steps = self.analysis.algorithm_steps[:4]
        num_steps = max(len(steps), 1)
        total_width = 11.7
        gap = 0.25
        card_w = (total_width - (gap * (num_steps - 1))) / num_steps
        step_top = 3.25
        step_h = 3.65

        for idx, step in enumerate(steps):
            s_left = 0.8 + idx * (card_w + gap)
            self._add_card(slide, left=s_left, top=step_top, width=card_w, height=step_h)

            tb_step = slide.shapes.add_textbox(
                Inches(s_left + 0.2), Inches(step_top + 0.2), Inches(card_w - 0.4), Inches(step_h - 0.4)
            )
            tf_s = tb_step.text_frame
            tf_s.word_wrap = True

            # Step Number Badge
            p_snum = tf_s.paragraphs[0]
            p_snum.text = f"PHASE {step.step_number or (idx + 1)}"
            p_snum.font.name = FONT_HEADING
            p_snum.font.size = Pt(10)
            p_snum.font.bold = True
            p_snum.font.color.rgb = COLOR_ACCENT_INDIGO
            p_snum.space_after = Pt(4)

            # Step Title
            p_stitle = tf_s.add_paragraph()
            p_stitle.text = step.title
            p_stitle.font.name = FONT_HEADING
            p_stitle.font.size = Pt(13)
            p_stitle.font.bold = True
            p_stitle.font.color.rgb = COLOR_TEXT_PRIMARY
            p_stitle.space_after = Pt(8)

            # Step Description
            p_sdesc = tf_s.add_paragraph()
            p_sdesc.text = step.description
            p_sdesc.font.name = FONT_BODY
            p_sdesc.font.size = Pt(11)
            p_sdesc.font.color.rgb = COLOR_TEXT_MUTED
            p_sdesc.line_spacing = 1.2

    # --------------------------------------------------------------------------
    # Slide 5: Benchmark Results & Empirical Evaluation (Light Theme)
    # --------------------------------------------------------------------------

    def build_benchmarks_slide(self):
        """Builds slide displaying empirical evaluation and benchmark comparisons."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_slide_background(slide, COLOR_LIGHT_BG)
        self._add_header(
            slide,
            kicker="04 // Empirical Evaluation",
            title="Benchmark Results & State-of-the-Art Comparison",
        )

        benchmarks: List[BenchmarkResult] = self.analysis.benchmark_results[:4]
        if not benchmarks:
            # Fallback placeholder if no benchmarks extracted
            benchmarks = [
                BenchmarkResult(
                    benchmark_name="Standard Systems Benchmark",
                    baseline="State of the Art Baseline",
                    result_metric="Significant Performance Gain",
                    significance="Demonstrates latency and throughput improvements over prior work.",
                )
            ]

        num_cards = len(benchmarks)
        if num_cards <= 2:
            # Two wide cards
            card_w = 5.7
            card_h = 5.0
            for idx, bm in enumerate(benchmarks):
                c_left = 0.8 + idx * 6.0
                self._add_card(slide, left=c_left, top=1.7, width=card_w, height=card_h)
                self._populate_benchmark_card(slide, bm, left=c_left, top=1.7, width=card_w, height=card_h)
        elif num_cards == 3:
            # Three column cards
            card_w = 3.7
            card_h = 5.1
            for idx, bm in enumerate(benchmarks):
                c_left = 0.8 + idx * 4.0
                self._add_card(slide, left=c_left, top=1.7, width=card_w, height=card_h)
                self._populate_benchmark_card(slide, bm, left=c_left, top=1.7, width=card_w, height=card_h)
        else:
            # 2x2 Grid for 4 benchmarks
            col_w = 5.7
            row_h = 2.45
            positions = [
                (0.8, 1.7),
                (6.8, 1.7),
                (0.8, 4.45),
                (6.8, 4.45),
            ]
            for idx, bm in enumerate(benchmarks[:4]):
                c_left, c_top = positions[idx]
                self._add_card(slide, left=c_left, top=c_top, width=col_w, height=row_h)
                self._populate_benchmark_card(slide, bm, left=c_left, top=c_top, width=col_w, height=row_h, compact=True)

    def _populate_benchmark_card(
        self, slide, bm: BenchmarkResult, left: float, top: float, width: float, height: float, compact: bool = False
    ):
        """Fills benchmark metric, baseline, and significance into a card."""
        tb = slide.shapes.add_textbox(Inches(left + 0.3), Inches(top + 0.25), Inches(width - 0.6), Inches(height - 0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        # Benchmark Name
        p_name = tf.paragraphs[0]
        p_name.text = bm.benchmark_name.upper()
        p_name.font.name = FONT_HEADING
        p_name.font.size = Pt(10)
        p_name.font.bold = True
        p_name.font.color.rgb = COLOR_TEXT_MUTED
        p_name.space_after = Pt(2)

        # Baseline Comparison Subtext
        p_base = tf.add_paragraph()
        p_base.text = f"Baseline: {bm.baseline}"
        p_base.font.name = FONT_BODY
        p_base.font.size = Pt(11)
        p_base.font.color.rgb = COLOR_TEXT_MUTED
        p_base.space_after = Pt(4)

        # Hero Metric (Vibrant Emerald Color)
        p_metric = tf.add_paragraph()
        p_metric.text = bm.result_metric
        p_metric.font.name = FONT_HEADING
        p_metric.font.size = Pt(18 if compact else 21)
        p_metric.font.bold = True
        p_metric.font.color.rgb = COLOR_ACCENT_GREEN
        p_metric.space_after = Pt(6)

        # Significance / Context Note
        p_sig = tf.add_paragraph()
        p_sig.text = bm.significance
        p_sig.font.name = FONT_BODY
        p_sig.font.size = Pt(10 if compact else 11)
        p_sig.font.color.rgb = COLOR_TEXT_PRIMARY
        p_sig.line_spacing = 1.15

    # --------------------------------------------------------------------------
    # Slide 6: System Implications & Takeaways (Light Theme)
    # --------------------------------------------------------------------------

    def build_conclusion_slide(self):
        """Builds slide summarizing real-world tradeoffs and conclusion takeaways."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_slide_background(slide, COLOR_LIGHT_BG)
        self._add_header(
            slide,
            kicker="05 // System Implications & Conclusion",
            title="Deployment Trade-Offs & Strategic Takeaways",
        )

        card_w = 5.65
        card_h = 5.2
        top_pos = 1.7

        # Left Card: System Implications & Trade-offs
        self._add_card(slide, left=0.8, top=top_pos, width=card_w, height=card_h)
        tb_l = slide.shapes.add_textbox(Inches(1.1), Inches(top_pos + 0.3), Inches(5.05), Inches(4.6))
        tf_l = tb_l.text_frame
        tf_l.word_wrap = True

        p_l_label = tf_l.paragraphs[0]
        p_l_label.text = "HARDWARE & ARCHITECTURAL TRADE-OFFS"
        p_l_label.font.name = FONT_HEADING
        p_l_label.font.size = Pt(11)
        p_l_label.font.bold = True
        p_l_label.font.color.rgb = COLOR_ACCENT_INDIGO
        p_l_label.space_after = Pt(10)

        p_l_text = tf_l.add_paragraph()
        p_l_text.text = self.analysis.system_tradeoffs_and_implications
        p_l_text.font.name = FONT_BODY
        p_l_text.font.size = Pt(13)
        p_l_text.font.color.rgb = COLOR_TEXT_PRIMARY
        p_l_text.line_spacing = 1.25

        # Right Card: Conclusion Takeaways
        self._add_card(slide, left=6.85, top=top_pos, width=card_w, height=card_h)
        tb_r = slide.shapes.add_textbox(Inches(7.15), Inches(top_pos + 0.3), Inches(5.05), Inches(4.6))
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True

        p_r_label = tf_r.paragraphs[0]
        p_r_label.text = "CORE TAKEAWAYS FOR PRACTITIONERS"
        p_r_label.font.name = FONT_HEADING
        p_r_label.font.size = Pt(11)
        p_r_label.font.bold = True
        p_r_label.font.color.rgb = COLOR_ACCENT_INDIGO
        p_r_label.space_after = Pt(10)

        for takeaway in self.analysis.conclusion_takeaways[:3]:
            p_item = tf_r.add_paragraph()
            p_item.text = f"✔  {takeaway}"
            p_item.font.name = FONT_BODY
            p_item.font.size = Pt(12)
            p_item.font.color.rgb = COLOR_TEXT_PRIMARY
            p_item.space_after = Pt(12)
            p_item.line_spacing = 1.25

    # --------------------------------------------------------------------------
    # Build Entire Presentation
    # --------------------------------------------------------------------------

    def build_deck(self, output_path: Union[str, Path]) -> Path:
        """
        Orchestrates creation of all slides and writes the presentation to disk.

        Args:
            output_path: Target path for the generated .pptx file.

        Returns:
            Resolved Path of the saved presentation.
        """
        logger.info("Generating Slide 1: Title & Executive Summary...")
        self.build_title_slide()

        logger.info("Generating Slide 2: Core Problem & Bottlenecks...")
        self.build_problem_slide()

        logger.info("Generating Slide 3: Key Technical Innovation...")
        self.build_innovation_slide()

        logger.info("Generating Slide 4: Algorithm & Execution Architecture...")
        self.build_algorithm_slide()

        logger.info("Generating Slide 5: Empirical Benchmark Results...")
        self.build_benchmarks_slide()

        logger.info("Generating Slide 6: System Implications & Takeaways...")
        self.build_conclusion_slide()

        dest = Path(output_path).resolve()
        dest.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(dest))
        logger.info(f"Presentation successfully saved to: {dest}")
        return dest


def generate_presentation(
    analysis: Union[PaperAnalysis, dict, str, Path],
    output_path: Union[str, Path] = "presentation.pptx",
) -> Path:
    """
    Generates a presentation deck from PaperAnalysis, JSON file, or dict.

    Args:
        analysis: PaperAnalysis instance, raw dict, or path to JSON file.
        output_path: Path for output .pptx presentation.

    Returns:
        Path to the saved PowerPoint file.
    """
    if isinstance(analysis, (str, Path)):
        json_file = Path(analysis)
        if json_file.is_file():
            data = json.loads(json_file.read_text(encoding="utf-8"))
            parsed_analysis = PaperAnalysis.model_validate(data)
        else:
            parsed_analysis = PaperAnalysis.model_validate_json(str(analysis))
    elif isinstance(analysis, dict):
        parsed_analysis = PaperAnalysis.model_validate(analysis)
    elif isinstance(analysis, PaperAnalysis):
        parsed_analysis = analysis
    else:
        raise TypeError(f"Unsupported analysis type: {type(analysis)}")

    builder = PresentationDeckBuilder(parsed_analysis)
    return builder.build_deck(output_path=output_path)


if __name__ == "__main__":
    import sys

    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")

    if len(sys.argv) < 2:
        print("Usage: python generator.py <path_to_metadata.json> [output.pptx]")
        sys.exit(1)

    input_json = sys.argv[1]
    out_pptx = sys.argv[2] if len(sys.argv) > 2 else "output_presentation.pptx"

    try:
        saved_file = generate_presentation(input_json, out_pptx)
        print(f"\nPresentation deck generated: {saved_file}")
    except Exception as err:
        print(f"Error generating presentation: {err}", file=sys.stderr)
        sys.exit(1)
